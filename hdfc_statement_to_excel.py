"""Single-file HDFC PDF-to-Excel converter with a desktop file chooser.

Run this file and a PDF selection window opens immediately.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import BinaryIO, Iterable

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader, PdfWriter


class StatementError(Exception):
    """A message that can safely be shown to the user."""


class PasswordRequiredError(StatementError):
    pass


class WrongPasswordError(StatementError):
    pass


@dataclass
class Transaction:
    date: datetime | None
    narration: str
    reference_no: str
    value_date: datetime | None
    withdrawal: Decimal | None
    deposit: Decimal | None
    closing_balance: Decimal | None
    page: int
    raw_date: str = ""
    raw_value_date: str = ""
    warnings: list[str] = field(default_factory=list)


@dataclass
class ConversionResult:
    transactions: list[Transaction]
    metadata: dict[str, str]
    warnings: list[str]
    workbook: bytes


_DATE_RE = re.compile(r"(?<!\d)([0-3]?\d)[/\-.]([01]?\d)[/\-.](\d{2}|\d{4})(?!\d)")


def _source_bytes(source: bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source
    source.seek(0)
    return source.read()


def unlock_pdf(source: bytes | BinaryIO, password: str = "") -> bytes:
    """Return a readable PDF, decrypting it when necessary."""
    data = _source_bytes(source)
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except Exception as exc:
        raise StatementError("The selected file is not a readable PDF.") from exc

    if not reader.is_encrypted:
        return data
    if not password:
        raise PasswordRequiredError("This PDF is password protected. Enter its password and try again.")
    try:
        result = reader.decrypt(password)
    except Exception as exc:
        raise WrongPasswordError("The PDF password is incorrect.") from exc
    if result == 0:
        raise WrongPasswordError("The PDF password is incorrect.")

    # pdfplumber handles most encrypted files directly, but writing a decrypted
    # in-memory copy gives consistent behavior across encryption algorithms.
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _ocr_digits(text: str) -> str:
    return text.translate(str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "|": "1"}))


def parse_date(text: str) -> datetime | None:
    cleaned = _ocr_digits(text).replace(" ", "")
    match = _DATE_RE.search(cleaned)
    parts: tuple[str, str, str] | None = match.groups() if match else None

    if parts is None:
        # Common OCR output loses slashes (13/05/25 -> 13105/25) or reads
        # them as a lower-case t/one (12/05/25 -> t2105t25). Reconstruct
        # only unmistakable dd/mm/yy shapes.
        fuzzy = cleaned.replace(",", "").replace(".", "").replace("-", "").replace("/", "")
        fuzzy = fuzzy.replace("t", "1")
        digits = re.sub(r"\D", "", fuzzy)
        if len(digits) == 6:
            parts = (digits[:2], digits[2:4], digits[4:])
        elif len(digits) == 7 and digits[2] == "1":
            parts = (digits[:2], digits[3:5], digits[5:])
        elif len(digits) == 8 and digits[2] == "1" and digits[5] == "1":
            parts = (digits[:2], digits[3:5], digits[6:])
    if parts is None:
        return None

    day, month, year = map(int, parts)
    if year < 100:
        year += 2000
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def parse_money(text: str) -> Decimal | None:
    if not text or not re.search(r"\d", text):
        return None
    cleaned = _ocr_digits(text).strip()
    # OCR can split the paise into a separate word: "19.500 00".
    cleaned = re.sub(r"\s+(\d{2})\s*$", r".\1", cleaned)
    cleaned = re.sub(r"^[rR](?=\d)", "1", cleaned)
    negative = cleaned.startswith("-") or ("(" in cleaned and ")" in cleaned)
    cleaned = re.sub(r"[^0-9.,]", "", cleaned)
    if not cleaned:
        return None

    # HDFC prints two decimal places. OCR frequently reads thousands separators
    # as dots (for example 40.000.00), so treat the final separator as decimal.
    separator_positions = [i for i, char in enumerate(cleaned) if char in ".,"]
    if separator_positions and len(cleaned) - separator_positions[-1] - 1 == 2:
        split = separator_positions[-1]
        whole = re.sub(r"\D", "", cleaned[:split]) or "0"
        fraction = re.sub(r"\D", "", cleaned[split + 1 :])
        canonical = f"{whole}.{fraction}"
    else:
        canonical = re.sub(r"\D", "", cleaned)
    try:
        value = Decimal(canonical)
        return -value if negative else value
    except InvalidOperation:
        return None


def _group_lines(words: Iterable[dict], tolerance: float = 4.0) -> list[list[dict]]:
    lines: list[list[dict]] = []
    for word in sorted(words, key=lambda w: (float(w["top"]), float(w["x0"]))):
        center = (float(word["top"]) + float(word["bottom"])) / 2
        for line in reversed(lines[-3:]):
            line_center = sum((float(w["top"]) + float(w["bottom"])) / 2 for w in line) / len(line)
            if abs(center - line_center) <= tolerance:
                line.append(word)
                break
        else:
            lines.append([word])
    return [sorted(line, key=lambda w: float(w["x0"])) for line in lines]


def _text_in(line: list[dict], left: float, right: float) -> str:
    return " ".join(w["text"] for w in line if float(w["x0"]) >= left and float(w["x0"]) < right).strip()


def _page_transactions(page, page_number: int) -> list[Transaction]:
    width = float(page.width)
    # Coordinates below are proportions of the standard HDFC A4 layout:
    # date, narration, reference, value date, withdrawal, deposit, balance.
    edges = [0.0, .112, .433, .538, .607, .727, .841, 1.01]
    edges = [edge * width for edge in edges]
    words = page.extract_words(x_tolerance=2, y_tolerance=3, keep_blank_chars=False)
    lines = _group_lines(words)
    rows: list[tuple[list[str], float]] = []
    for line in lines:
        top = min(float(w["top"]) for w in line)
        if top < page.height * .28 or top > page.height * .92:
            continue
        cells = [_text_in(line, edges[i], edges[i + 1]) for i in range(7)]
        # A transaction's first line has text in the date column and at least
        # one amount at the right. This also works when OCR damages the date.
        right_values = [parse_money(cells[i]) for i in (4, 5, 6)]
        is_start = bool(cells[0]) and any(value is not None for value in right_values)
        if is_start:
            rows.append((cells, top))
        elif rows and cells[1] and not any(cells[i] for i in (0, 4, 5, 6)):
            # Narration/reference continuation line.
            previous = rows[-1][0]
            previous[1] = " ".join(filter(None, (previous[1], cells[1])))
            if cells[2]:
                previous[2] = " ".join(filter(None, (previous[2], cells[2])))

    transactions: list[Transaction] = []
    for cells, _ in rows:
        raw_date, narration, reference, raw_value, raw_withdrawal, raw_deposit, raw_balance = cells
        if "STATEMENT SUMMARY" in narration.upper():
            continue
        date = parse_date(raw_date)
        value_date = parse_date(raw_value)
        warnings: list[str] = []
        if not date:
            warnings.append(f"Could not read transaction date: {raw_date!r}; value date used")
            date = value_date
        if raw_value and not value_date:
            warnings.append(f"Could not read value date: {raw_value!r}; transaction date used")
            value_date = date
        # Booking and value dates can differ by days, not a decade. Correct a
        # clearly damaged two-digit year from the other date.
        if date and value_date and abs(date.year - value_date.year) > 1:
            try:
                date = date.replace(year=value_date.year)
                warnings.append("Transaction year corrected from value date")
            except ValueError:
                pass
        withdrawal = parse_money(raw_withdrawal)
        deposit = parse_money(raw_deposit)
        balance = parse_money(raw_balance)
        if raw_balance and balance is None:
            warnings.append(f"Could not read closing balance: {raw_balance!r}")
        transactions.append(Transaction(
            date=date,
            narration=re.sub(r"\s+", " ", narration).strip(),
            reference_no=re.sub(r"\s+", "", reference),
            value_date=value_date,
            withdrawal=withdrawal,
            deposit=deposit,
            closing_balance=balance,
            page=page_number,
            raw_date=raw_date,
            raw_value_date=raw_value,
            warnings=warnings,
        ))
    return transactions


def _metadata(text: str) -> dict[str, str]:
    compact = re.sub(r"[ \t]+", " ", text)
    fields: dict[str, str] = {}
    patterns = {
        "Account number": r"Account No\s*:?\s*([0-9OIl]{10,18})",
        "Customer ID": r"Cust(?:omer)?\s*ID\s*:?\s*([0-9OIl]{6,15})",
        "IFSC": r"(?:RTGS/NEFT\s*)?IFSC\s*:?\s*([A-Z]{4}0[A-Z0-9]{6})",
        "Statement period": r"From\s*:?\s*([0-9/\-.]+)\s+To\s*:?\s*([0-9/\-.]+)",
    }
    for label, pattern in patterns.items():
        match = re.search(pattern, compact, flags=re.IGNORECASE)
        if match:
            value = " to ".join(match.groups())
            fields[label] = _ocr_digits(value) if label != "Statement period" else value
    return fields


def create_workbook(transactions: list[Transaction], metadata: dict[str, str], warnings: list[str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    headers = ["Date", "Narration", "Reference No.", "Value Date", "Withdrawal", "Deposit", "Closing Balance", "PDF Page", "Review Notes"]
    sheet.append(headers)
    for transaction in transactions:
        sheet.append([
            transaction.date,
            transaction.narration,
            transaction.reference_no,
            transaction.value_date,
            transaction.withdrawal,
            transaction.deposit,
            transaction.closing_balance,
            transaction.page,
            "; ".join(transaction.warnings),
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "dd-mmm-yyyy"
        row[3].number_format = "dd-mmm-yyyy"
        for index in (4, 5, 6):
            row[index].number_format = '#,##0.00;[Red]-#,##0.00'
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[8].alignment = Alignment(wrap_text=True, vertical="top")
    widths = [14, 62, 25, 14, 16, 16, 18, 10, 45]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    if transactions:
        table = Table(displayName="HDFCTransactions", ref=f"A1:I{len(transactions) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)

    info = workbook.create_sheet("Statement Info")
    info.append(["Field", "Value"])
    for key, value in metadata.items():
        info.append([key, value])
    info.append(["Transactions extracted", len(transactions)])
    info.append(["Rows needing review", sum(bool(t.warnings) for t in transactions)])
    if warnings:
        info.append(["General warnings", " | ".join(warnings)])
    for cell in info[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 90
    info.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def convert_statement(source: bytes | BinaryIO, password: str = "") -> ConversionResult:
    pdf_data = unlock_pdf(source, password)
    transactions: list[Transaction] = []
    text_parts: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(pdf_data)) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                text_parts.append(page.extract_text() or "")
                transactions.extend(_page_transactions(page, page_number))
    except Exception as exc:
        raise StatementError("The PDF could not be read. It may be damaged or use an unsupported format.") from exc

    if not transactions:
        raise StatementError(
            "No HDFC transaction rows were found. This tool expects HDFC's standard statement layout and a PDF containing selectable/OCR text."
        )
    warnings: list[str] = []
    review_count = sum(bool(t.warnings) for t in transactions)
    if review_count:
        warnings.append(f"{review_count} row(s) contain OCR fields that should be reviewed in Excel.")
    metadata = _metadata("\n".join(text_parts))
    workbook = create_workbook(transactions, metadata, warnings)
    return ConversionResult(transactions, metadata, warnings, workbook)

# ---------------------------------------------------------------------------
# Desktop file-selection interface
# ---------------------------------------------------------------------------

def run_desktop_converter() -> None:
    """Open a PDF chooser immediately, then save the converted Excel file."""
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except ImportError as exc:
        raise SystemExit(
            "Tkinter is required for the file-selection window. On Ubuntu/Debian, "
            "install it with: sudo apt install python3-tk"
        ) from exc

    root = tk.Tk()
    root.withdraw()
    root.update()

    pdf_path = filedialog.askopenfilename(
        parent=root,
        title="Select HDFC Bank statement PDF",
        filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
    )
    if not pdf_path:
        root.destroy()
        return

    try:
        pdf_data = open(pdf_path, "rb").read()
        reader = PdfReader(io.BytesIO(pdf_data), strict=False)
    except Exception:
        messagebox.showerror("Invalid PDF", "The selected file is not a readable PDF.", parent=root)
        root.destroy()
        return

    password = ""
    if reader.is_encrypted:
        password = simpledialog.askstring(
            "PDF password",
            "This statement is password protected.\nEnter the PDF password:",
            show="*",
            parent=root,
        )
        if password is None:
            root.destroy()
            return

    try:
        result = convert_statement(pdf_data, password)
    except WrongPasswordError:
        messagebox.showerror("Incorrect password", "The PDF password is incorrect.", parent=root)
        root.destroy()
        return
    except StatementError as exc:
        messagebox.showerror("Conversion failed", str(exc), parent=root)
        root.destroy()
        return
    except Exception as exc:
        messagebox.showerror("Conversion failed", f"Unexpected error: {exc}", parent=root)
        root.destroy()
        return

    from pathlib import Path

    source = Path(pdf_path)
    output_path = filedialog.asksaveasfilename(
        parent=root,
        title="Save Excel file",
        initialdir=str(source.parent),
        initialfile=f"{source.stem}_transactions.xlsx",
        defaultextension=".xlsx",
        filetypes=[("Excel workbook", "*.xlsx")],
    )
    if not output_path:
        root.destroy()
        return

    try:
        Path(output_path).write_bytes(result.workbook)
    except OSError as exc:
        messagebox.showerror("Could not save file", str(exc), parent=root)
        root.destroy()
        return

    review_count = sum(bool(row.warnings) for row in result.transactions)
    review_message = (
        f"\n\n{review_count} row(s) were marked for OCR review in the Review Notes column."
        if review_count
        else ""
    )
    messagebox.showinfo(
        "Conversion complete",
        f"Created {len(result.transactions)} transaction rows.\n\nSaved to:\n{output_path}{review_message}",
        parent=root,
    )
    root.destroy()


if __name__ == "__main__":
    run_desktop_converter()
